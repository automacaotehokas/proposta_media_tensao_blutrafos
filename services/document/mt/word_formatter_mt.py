from docx.shared import Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_row_height(row, height_cm: float) -> None:
    """
    Define a altura de uma linha da tabela.
    
    Args:
        row: Linha da tabela
        height_cm: Altura em centímetros
    """
    tr = row._tr
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(int(height_cm * 567)))  # Multiplicando por 567 para converter cm para twips
    trHeight.set(qn('w:hRule'), 'exact')  # Define altura exata
    tr.append(trHeight)

def set_column_widths(table, col_widths: list) -> None:
    """
    Define a largura das colunas da tabela.
    
    Args:
        table: Tabela do documento
        col_widths: Lista com as larguras das colunas
    """
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

def apply_paragraph_formatting(paragraph, alignment: str = 'center', 
                          space_before: Pt = Pt(5), space_after: Pt = Pt(0)) -> None:
    """
    Aplica formatação a um parágrafo.
    
    Args:
        paragraph: Parágrafo a ser formatado
        alignment: Alinhamento ('left', 'center', 'right')
        space_before: Espaço antes do parágrafo
        space_after: Espaço depois do parágrafo
    """
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_before = space_before
    paragraph_format.space_after = space_after
    paragraph_format.line_spacing = 1  # Espaçamento simples
    paragraph.alignment = {'left': 0, 'center': 1, 'right': 2}.get(alignment, 1)

def set_cell_shading(cell, color: str) -> None:
    """
    Aplica sombreamento a uma célula.
    
    Args:
        cell: Célula da tabela
        color: Cor em formato hexadecimal (ex: '00543C')
    """
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)

def add_double_borders(cell) -> None:
    """
    Adiciona bordas duplas a uma célula.
    
    Args:
        cell: Célula da tabela
    """
    tc = cell._element
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    # Adiciona borda para cada lado
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'double')  # Bordas duplas
        border.set(qn('w:sz'), '4')        # Espessura
        border.set(qn('w:space'), '0')     # Sem espaço
        border.set(qn('w:color'), '000000') # Cor preta
        tcBorders.append(border)
    
    tcPr.append(tcBorders)

def set_table_left_indent(table, indent: int) -> None:
    """
    Define a indentação esquerda da tabela.
    
    Args:
        table: Tabela do documento
        indent: Valor da indentação em twips
    """
    tbl_pr = table._tbl.tblPr
    tbl_indent = tbl_pr.xpath("w:tblInd")
    
    if tbl_indent:
        tbl_indent[0].set(qn('w:w'), str(indent))
        tbl_indent[0].set(qn('w:type'), 'dxa')
    else:
        tbl_indent_element = OxmlElement('w:tblInd')
        tbl_indent_element.set(qn('w:w'), str(indent))
        tbl_indent_element.set(qn('w:type'), 'dxa')
        tbl_pr.append(tbl_indent_element)

def apply_cell_formatting(cell, font_name: str = 'Calibri Light (Títulos)', 
                        font_size: int = 11, bold: bool = False, 
                        alignment: str = 'center', color: str = None) -> None:
    """
    Aplica formatação completa a uma célula.
    
    Args:
        cell: Célula da tabela
        font_name: Nome da fonte
        font_size: Tamanho da fonte
        bold: Se a fonte deve ser negrito
        alignment: Alinhamento do texto
        color: Cor do texto (opcional)
    """
    paragraph = cell.paragraphs[0]
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    
    apply_paragraph_formatting(paragraph, alignment=alignment)

def merge_cells_range(table, start_row: int, start_col: int, 
                     end_row: int, end_col: int) -> None:
    """
    Mescla um intervalo de células em uma tabela.
    
    Args:
        table: Tabela do documento
        start_row: Linha inicial
        start_col: Coluna inicial
        end_row: Linha final
        end_col: Coluna final
    """
    cell_range = table.rows[start_row].cells[start_col:end_col + 1]
    if cell_range:
        cell_range[0].merge(table.rows[end_row].cells[end_col])

def apply_table_style(table, header_color: str = '00543C', 
                     font_name: str = 'Calibri Light (Títulos)', 
                     font_size: int = 11) -> None:
    """
    Aplica estilo padrão a uma tabela inteira.
    
    Args:
        table: Tabela do documento
        header_color: Cor do cabeçalho
        font_name: Nome da fonte padrão
        font_size: Tamanho da fonte padrão
    """
    # Formata o cabeçalho
    for cell in table.rows[0].cells:
        apply_cell_formatting(cell, font_name, font_size, True, 'center', 'FFFFFF')
        set_cell_shading(cell, header_color)
        add_double_borders(cell)
    
    # Formata as células do corpo
    for row in table.rows[1:]:
        for cell in row.cells:
            apply_cell_formatting(cell, font_name, font_size)
            add_double_borders(cell)